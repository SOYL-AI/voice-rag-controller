"""
make_recording_kit.py: a small real-speech check (exploratory, post hoc).

Selects 120 d1_open TEST items (alias-graded, so no LLM judge is needed),
balanced by empirical need and item type, with questions that are easy to
read aloud (8-22 words, no digits-heavy strings), and splits them into one
sheet of 20 per author. Writes, into --out-dir:
  Recording_<Name>.xlsx      the questions to read, one per row, with file names
  READ_ME_RECORDING.txt      instructions
  recording_master.csv       item ids, need labels and assignment (keep private)

    python analysis/make_recording_kit.py --out-dir <folder>
"""

import argparse
import csv
import json
import random
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
SEED = 20261019
AUTHORS = ["Renya", "Ryan", "Shubhankit", "Lakshya", "Adithi", "Yashmita"]
PER_AUTHOR = 20

INSTRUCTIONS = """REAL-SPEECH RECORDING (Paper 4422) - about 15 minutes per person

What: read 20 short questions aloud, one audio file per question.

1. Use your phone's voice recorder (or any recorder). A quiet room is best,
   but normal background noise is fine. Do not use earphones with a mic that
   rubs on clothing.
2. For each row of your sheet: start recording, wait about one second, read
   the question once at your normal speaking speed, wait about one second,
   stop. Read it the way you would ask a voice assistant. Do not rush and do
   not over-pronounce.
3. If you stumble, just record that question again (keep only the good one).
4. Name each file exactly as in the "File name" column (for example
   R03_Renya.m4a). Any format is fine: m4a, mp3, wav, ogg.
5. Put all 20 files in one folder named with your name and share it (Drive or
   WhatsApp as files, not as voice notes if possible, since voice notes are
   compressed heavily).

Please do not look up the answers or change the wording of the questions.
"""


def readable(q: str) -> bool:
    words = q.split()
    if not 8 <= len(words) <= 22:
        return False
    if len(re.findall(r"\d", q)) > 4:
        return False
    if re.search(r"[\[\]{}<>|_#@^~\\]", q):
        return False
    return True


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", type=Path, required=True)
    args = ap.parse_args()

    items = {}
    for line in open(REPO / "datasets/v2/d1_open/items.jsonl", encoding="utf-8"):
        it = json.loads(line)
        if it.get("split") == "test":
            items[it["id"]] = it
    labels = json.load(open(REPO / "llm/experiment1_v2/results/d1_open/mistral7b_v03/labels_test.json",
                            encoding="utf-8"))["items"]

    strata = {}
    for iid, it in items.items():
        lab = labels.get(iid)
        if lab is None or lab.get("empirical_need") is None or not readable(it["question"]):
            continue
        key = (it["item_type"], bool(lab["empirical_need"]))
        strata.setdefault(key, []).append(iid)

    n_total = PER_AUTHOR * len(AUTHORS)
    # multihop is rare (213 of 3,734); take 20 of it, the rest PopQA, need balanced 50/50
    quota = {("multihop", True): 10, ("multihop", False): 10,
             ("popqa", True): (n_total - 20) // 2, ("popqa", False): (n_total - 20) // 2}
    rng = random.Random(SEED)
    chosen = []
    for key, q in quota.items():
        pool = sorted(strata.get(key, []))
        rng.shuffle(pool)
        if len(pool) < q:
            raise SystemExit(f"stratum {key} has only {len(pool)} readable items, need {q}")
        chosen += [(iid, key) for iid in pool[:q]]
    rng.shuffle(chosen)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, (iid, key) in enumerate(chosen):
        author = AUTHORS[i // PER_AUTHOR]
        n = i % PER_AUTHOR + 1
        rows.append({"author": author, "row": n, "file_name": f"R{n:02d}_{author}",
                     "item_id": iid, "item_type": key[0], "empirical_need": key[1],
                     "question": items[iid]["question"]})

    with open(args.out_dir / "recording_master.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    (args.out_dir / "READ_ME_RECORDING.txt").write_text(INSTRUCTIONS, encoding="utf-8")

    head_fill = PatternFill("solid", fgColor="DDE8F6")
    for author in AUTHORS:
        wb = Workbook()
        ws = wb.active
        ws.title = "Questions"
        ws.append(["Row", "File name (use exactly)", "Question to read aloud", "Recorded? (tick)"])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for r in rows:
            if r["author"] == author:
                ws.append([r["row"], r["file_name"], r["question"], ""])
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 90
        ws.column_dimensions["D"].width = 16
        for row in ws.iter_rows(min_row=2):
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
        ws2 = wb.create_sheet("Instructions")
        for line in INSTRUCTIONS.splitlines():
            ws2.append([line])
        ws2.column_dimensions["A"].width = 100
        wb.save(args.out_dir / f"Recording_{author}.xlsx")
    print(f"wrote {len(rows)} items for {len(AUTHORS)} authors to {args.out_dir}")


if __name__ == "__main__":
    main()
